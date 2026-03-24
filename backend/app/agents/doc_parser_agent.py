"""
文档解析Agent

支持多格式文档解析：PDF, DOCX, DOC, TXT
"""

import os
from typing import Any

from .base_agent import BaseAgent


class DocParserAgent(BaseAgent):
    """文档解析Agent"""

    def __init__(self, llm_client=None):
        super().__init__("doc_parser", llm_client)

    async def execute(self, input_data: dict) -> Any:
        filepath = input_data.get("filepath", "")
        return self.parse(filepath)

    def parse(self, filepath: str) -> dict:
        """解析文档，返回文本内容"""
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"文件不存在: {filepath}")

        ext = os.path.splitext(filepath)[1].lower()
        parsers = {
            ".pdf": self._parse_pdf,
            ".docx": self._parse_docx,
            ".doc": self._parse_doc,
            ".txt": self._parse_txt,
        }

        parser = parsers.get(ext)
        if not parser:
            raise ValueError(f"不支持的文件类型: {ext}")

        text = parser(filepath)
        return {
            "text": text,
            "file_type": ext,
            "char_count": len(text),
            "filename": os.path.basename(filepath),
        }

    def _parse_pdf(self, filepath: str) -> str:
        """解析PDF文件"""
        import fitz
        doc = fitz.open(filepath)
        text = "".join([page.get_text() for page in doc])
        doc.close()
        return text

    def _parse_docx(self, filepath: str) -> str:
        """解析DOCX文件"""
        from docx import Document
        doc = Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n\n".join(paragraphs)

    def _parse_doc(self, filepath: str) -> str:
        """解析DOC文件（旧版Word二进制格式）

        优先用 antiword 提取文本，若不可用则通过 LibreOffice 转换为 docx 后解析。
        """
        import subprocess
        import tempfile

        # 方式1: antiword（轻量，速度快）
        try:
            result = subprocess.run(
                ["antiword", filepath],
                capture_output=True, text=True, timeout=30,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout.strip()
        except FileNotFoundError:
            pass  # antiword 未安装，尝试下一种方式

        # 方式2: LibreOffice 转 docx 后用 python-docx 解析
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                subprocess.run(
                    ["libreoffice", "--headless", "--convert-to", "docx",
                     "--outdir", tmpdir, filepath],
                    capture_output=True, timeout=60,
                    check=True,
                )
                basename = os.path.splitext(os.path.basename(filepath))[0]
                converted = os.path.join(tmpdir, f"{basename}.docx")
                if os.path.exists(converted):
                    return self._parse_docx(converted)
                raise RuntimeError("LibreOffice 转换后未生成 docx 文件")
        except FileNotFoundError:
            raise RuntimeError(
                "解析 .doc 文件需要安装 antiword 或 LibreOffice。\n"
                "  Ubuntu/Debian: apt-get install antiword\n"
                "  或: apt-get install libreoffice"
            )

    def _parse_txt(self, filepath: str) -> str:
        """解析纯文本文件"""
        with open(filepath, "r", encoding="utf-8") as f:
            return f.read()
